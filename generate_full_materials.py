import os
import urllib.request
from docx import Document
from fpdf import FPDF

os.makedirs("data/kb", exist_ok=True)

print("Генерация TXT файла (Курс по SQL)...")
# --- 1. TXT: 3 полноценных урока по SQL ---
txt_content = """КУРС: БАЗЫ ДАННЫХ И SQL (ПОЛНОЦЕННОЕ РУКОВОДСТВО)

================================================================================
УРОК 1: ВВЕДЕНИЕ В РЕЛЯЦИОННЫЕ БАЗЫ ДАННЫХ И БАЗОВЫЙ SQL
================================================================================
Реляционные базы данных (РБД) хранят информацию в виде связанных таблиц, состоящих из строк (записей) и столбцов (атрибутов). Самый популярный язык для работы с РБД — SQL (Structured Query Language).

1.1. Базовая выборка данных (SELECT)
Команда SELECT используется для чтения данных из таблиц.
- Выбрать все колонки: 
  SELECT * FROM employees;
- Выбрать конкретные колонки: 
  SELECT first_name, last_name, salary FROM employees;

1.2. Фильтрация данных (WHERE)
Чтобы получить не все записи, а только те, которые соответствуют условию, используется WHERE.
- Операторы сравнения: =, <, >, <=, >=, != (или <>)
  SELECT * FROM employees WHERE salary > 50000;
- Логические операторы (AND, OR, NOT):
  SELECT * FROM employees WHERE department = 'IT' AND salary > 50000;
- Оператор IN (поиск по списку значений):
  SELECT * FROM employees WHERE department IN ('IT', 'HR', 'Sales');
- Оператор LIKE (поиск по шаблону, % заменяет любое количество символов):
  SELECT * FROM employees WHERE last_name LIKE 'Smi%';

1.3. Сортировка и ограничение (ORDER BY, LIMIT)
- Сортировка по убыванию (DESC) или возрастанию (ASC, по умолчанию):
  SELECT * FROM employees ORDER BY salary DESC;
- Ограничение количества выводимых строк (полезно для пагинации):
  SELECT * FROM employees ORDER BY salary DESC LIMIT 10;


================================================================================
УРОК 2: ПРОДВИНУТАЯ ВЫБОРКА ДАННЫХ (JOIN И АГРЕГАЦИЯ)
================================================================================
В реальных проектах данные разнесены по разным таблицам (нормализация). Чтобы собрать их вместе, используются объединения (JOIN).

2.1. Объединение таблиц (JOIN)
Представим две таблицы: `users` (id, name, role_id) и `roles` (id, role_name).
- INNER JOIN (внутреннее соединение) — возвращает только те строки, для которых есть совпадения в обеих таблицах:
  SELECT users.name, roles.role_name 
  FROM users 
  INNER JOIN roles ON users.role_id = roles.id;
- LEFT JOIN (левое соединение) — возвращает ВСЕ строки из левой таблицы (users), даже если в правой (roles) нет совпадений (вместо них будет NULL):
  SELECT users.name, roles.role_name 
  FROM users 
  LEFT JOIN roles ON users.role_id = roles.id;

2.2. Агрегатные функции
Позволяют вычислять сводные значения по множеству строк.
- COUNT() — количество строк.
- SUM() — сумма значений.
- AVG() — среднее значение.
- MIN() / MAX() — минимальное и максимальное значение.
Пример: SELECT COUNT(*), AVG(salary) FROM employees;

2.3. Группировка данных (GROUP BY)
Используется вместе с агрегатными функциями для группировки результата по одному или нескольким столбцам.
Пример (средняя зарплата по каждому отделу):
  SELECT department, AVG(salary) 
  FROM employees 
  GROUP BY department;

2.4. Фильтрация групп (HAVING)
WHERE фильтрует строки ДО группировки, а HAVING — ПОСЛЕ.
Пример (отделы, где средняя зарплата больше 60000):
  SELECT department, AVG(salary) 
  FROM employees 
  GROUP BY department 
  HAVING AVG(salary) > 60000;


================================================================================
УРОК 3: МОДИФИКАЦИЯ ДАННЫХ И ПРОЕКТИРОВАНИЕ СХЕМЫ (DDL/DML)
================================================================================
SQL делится на подмножества. DML (Data Manipulation Language) работает с данными, а DDL (Data Definition Language) — со структурой таблиц.

3.1. Создание таблиц (CREATE TABLE)
При создании таблицы нужно указать типы данных (INT, VARCHAR, DATE, BOOLEAN) и ограничения (CONSTRAINTS).
- PRIMARY KEY — уникальный идентификатор строки.
- FOREIGN KEY — ссылка на первичный ключ другой таблицы.
- NOT NULL — поле не может быть пустым.
- UNIQUE — все значения в столбце должны быть уникальными.

Пример:
CREATE TABLE orders (
    id INT PRIMARY KEY AUTO_INCREMENT,
    user_id INT NOT NULL,
    total_amount DECIMAL(10, 2),
    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
    FOREIGN KEY (user_id) REFERENCES users(id)
);

3.2. Вставка данных (INSERT)
INSERT INTO users (name, email, age) 
VALUES ('Иван Иванов', 'ivan@example.com', 28);

3.3. Обновление данных (UPDATE)
ВНИМАНИЕ: Всегда используйте WHERE при UPDATE, иначе обновятся ВСЕ строки в таблице!
UPDATE users 
SET age = 29, email = 'new_ivan@example.com' 
WHERE id = 1;

3.4. Удаление данных (DELETE)
ВНИМАНИЕ: Аналогично UPDATE, всегда используйте WHERE. В современных системах часто используют "мягкое удаление" (soft delete), просто меняя статус (UPDATE users SET is_deleted = true WHERE id = 1).
DELETE FROM users WHERE id = 1;

3.5. Транзакции (ACID)
Транзакция — это набор операций, которые выполняются как единое целое. Либо выполняются все, либо ни одной.
BEGIN; -- начало транзакции
UPDATE accounts SET balance = balance - 100 WHERE id = 1;
UPDATE accounts SET balance = balance + 100 WHERE id = 2;
COMMIT; -- успешное завершение и сохранение
-- Если произошла ошибка, используем ROLLBACK; для отмены изменений.
"""
with open("data/kb/course_sql_full.txt", "w", encoding="utf-8") as f:
    f.write(txt_content)


print("Генерация DOCX файла (Курс по FastAPI)...")
# --- 2. DOCX: 3 полноценных урока по FastAPI ---
doc = Document()
doc.add_heading('Курс: Современная веб-разработка на FastAPI', 0)

# Урок 1
doc.add_heading('Урок 1: Введение в FastAPI и создание первого API', level=1)
doc.add_paragraph('FastAPI — это современный, быстрый (высокопроизводительный) веб-фреймворк для создания API на Python 3.7+, в основе которого лежат стандартные аннотации типов Python (type hints).')

doc.add_heading('1.1. Установка и запуск', level=2)
doc.add_paragraph('Для работы нам потребуется сам фреймворк и ASGI-сервер (обычно используется uvicorn).')
doc.add_paragraph('Команда для установки:\npip install fastapi uvicorn')

doc.add_heading('1.2. Создание первого эндпоинта', level=2)
doc.add_paragraph('Создайте файл main.py и напишите следующий код:')
doc.add_paragraph('from fastapi import FastAPI\n\napp = FastAPI()\n\n@app.get("/")\ndef read_root():\n    return {"message": "Hello World"}')
doc.add_paragraph('В этом коде мы создаем экземпляр приложения FastAPI и используем декоратор @app.get("/"), чтобы указать, что функция read_root будет обрабатывать GET-запросы по корневому URL.')

doc.add_heading('1.3. Запуск сервера', level=2)
doc.add_paragraph('Запустите сервер командой:\nuvicorn main:app --reload')
doc.add_paragraph('Флаг --reload означает, что сервер будет автоматически перезапускаться при изменении кода (полезно для разработки). FastAPI автоматически генерирует интерактивную документацию Swagger, доступную по адресу http://127.0.0.1:8000/docs.')

# Урок 2
doc.add_heading('Урок 2: Обработка параметров и валидация данных (Pydantic)', level=1)

doc.add_heading('2.1. Path-параметры (Параметры пути)', level=2)
doc.add_paragraph('Параметры пути используются для идентификации конкретного ресурса.')
doc.add_paragraph('@app.get("/users/{user_id}")\ndef get_user(user_id: int):\n    return {"user_id": user_id}')
doc.add_paragraph('FastAPI автоматически прочитает user_id из URL и преобразует его в целое число (int). Если передать строку (например, /users/abc), FastAPI автоматически вернет красивую JSON-ошибку валидации.')

doc.add_heading('2.2. Query-параметры (Параметры запроса)', level=2)
doc.add_paragraph('Если параметр функции не указан в пути (URL), FastAPI интерпретирует его как query-параметр (то, что идет после знака ? в URL).')
doc.add_paragraph('@app.get("/items/")\ndef read_items(skip: int = 0, limit: int = 10):\n    return {"skip": skip, "limit": limit}')

doc.add_heading('2.3. Тела запросов (Request Body) и Pydantic', level=2)
doc.add_paragraph('Для получения данных от клиента (например, при регистрации пользователя) используются POST-запросы и модели Pydantic.')
doc.add_paragraph('from pydantic import BaseModel\n\nclass Item(BaseModel):\n    name: str\n    description: str | None = None\n    price: float\n\n@app.post("/items/")\ndef create_item(item: Item):\n    return {"item_name": item.name, "price_with_tax": item.price * 1.2}')
doc.add_paragraph('Pydantic автоматически проверяет, что клиент прислал JSON с нужными полями правильных типов.')

# Урок 3
doc.add_heading('Урок 3: Обработка ошибок и Dependency Injection', level=1)

doc.add_heading('3.1. Возврат HTTP-ошибок', level=2)
doc.add_paragraph('Если ресурс не найден или у пользователя нет прав, нужно вернуть соответствующий HTTP-статус. В FastAPI для этого используется HTTPException.')
doc.add_paragraph('from fastapi import HTTPException\n\n@app.get("/items/{item_id}")\ndef get_item(item_id: int):\n    if item_id not in items_db:\n        raise HTTPException(status_code=404, detail="Item not found")\n    return items_db[item_id]')

doc.add_heading('3.2. Внедрение зависимостей (Dependency Injection)', level=2)
doc.add_paragraph('DI — это мощный механизм FastAPI, позволяющий переиспользовать логику (например, подключение к БД, проверка токена авторизации) без дублирования кода.')
doc.add_paragraph('from fastapi import Depends\n\ndef verify_token(token: str):\n    if token != "supersecret":\n        raise HTTPException(status_code=401, detail="Invalid token")\n    return True\n\n@app.get("/secure-data/")\ndef get_secure_data(is_valid: bool = Depends(verify_token)):\n    return {"data": "This is highly confidential"}')

doc.save("data/kb/course_fastapi_full.docx")


print("Генерация PDF файла (Best Practices)...")
# --- 3. PDF: Памятка по Best Practices ---
font_path = "FreeSans.ttf"
if not os.path.exists(font_path):
    print("Скачивание шрифта FreeSans для поддержки кириллицы в PDF...")
    urllib.request.urlretrieve("https://github.com/opensourcedesign/fonts/raw/master/gnu-freefont_freesans/FreeSans.ttf", font_path)

pdf = FPDF()
pdf.add_page()
pdf.add_font("FreeSans", style="", fname=font_path)
pdf.set_font("FreeSans", size=12)

pdf_content = """ПАМЯТКА: ЧИСТЫЙ КОД И BEST PRACTICES В PYTHON

1. СТАНДАРТ PEP 8 (ОФОРМЛЕНИЕ КОДА)
- Отступы: Строго 4 пробела. Никогда не используйте табуляцию (Tab).
- Длина строки: Ограничивайте строки 79 или 120 символами. Это улучшает читаемость при разделении экрана на две части.
- Именование переменных и функций: Используйте snake_case (например: get_user_data, total_amount).
- Именование классов: Используйте CamelCase (например: UserProfile, DatabaseConnection).
- Именование констант: Используйте UPPER_SNAKE_CASE (например: MAX_RETRIES = 5).
- Импорты: Все импорты должны быть в самом начале файла. Сначала стандартные библиотеки, затем сторонние (pip), затем локальные модули проекта.

2. ТАЙП-ХИНТЫ (АННОТАЦИИ ТИПОВ)
Python — язык с динамической типизацией, но использование аннотаций типов (Type Hints) стало индустриальным стандартом. Они помогают IDE (VS Code, PyCharm) находить ошибки до запуска кода.

ПЛОХО:
def calculate_discount(price, discount):
    return price - (price * discount)

ХОРОШО:
def calculate_discount(price: float, discount: float) -> float:
    return price - (price * discount)

3. ЧАСТЫЕ ОШИБКИ И АНТИПАТТЕРНЫ (ЧЕГО ДЕЛАТЬ НЕЛЬЗЯ)

Антипаттерн 1: Изменяемые объекты по умолчанию
ПЛОХО: def add_item(item, my_list=[]): ...
Проблема: Список создастся один раз при инициализации функции и будет сохранять состояние между вызовами.
ХОРОШО: 
def add_item(item, my_list=None):
    if my_list is None:
        my_list = []
    my_list.append(item)

Антипаттерн 2: "Голый" except
ПЛОХО:
try:
    x = 1 / 0
except:
    pass
Проблема: Это перехватит абсолютно все ошибки, включая системные (например, прерывание с клавиатуры Ctrl+C) и опечатки в коде.
ХОРОШО:
try:
    x = 1 / 0
except ZeroDivisionError as e:
    print(f"Ошибка деления: {e}")

4. ПРИНЦИПЫ SOLID И DRY
- DRY (Don't Repeat Yourself): Не дублируйте код. Если вы скопировали блок кода три раза — вынесите его в отдельную функцию.
- Single Responsibility (Единая ответственность): Одна функция или класс должны решать только одну задачу. Если функция называется get_data_and_save_to_db_and_send_email(), её нужно разбить на три разные функции.
"""

# Разбиваем текст на строки и добавляем в PDF
for line in pdf_content.split('\n'):
    pdf.multi_cell(190, 8, text=line)

pdf.output("data/kb/python_best_practices.pdf")

print("Все файлы успешно сгенерированы в папке data/kb!")
